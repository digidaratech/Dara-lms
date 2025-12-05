import os
import logging
import re
from typing import List

# External libraries
from docx import Document  # python-docx
from pptx import Presentation  # python-pptx
from pdfminer.high_level import extract_text   # pdfminer.six

logger = logging.getLogger(__name__)

class DocumentProcessor:
    def __init__(self, base_courses_dir: str = "courses"):
        self.supported_extensions = ['.txt', '.pdf', '.docx', '.pptx']
        self.chunk_size = 450    # Optimized chunk size
        self.overlap = 80        # Better continuity

        BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        self.base_courses_dir = os.path.join(BASE_DIR, "courses")
        logger.info(f"[DOC] Base path set: {self.base_courses_dir}")

    # Extract Text from Files
    def extract_text_from_file(self, file_path: str) -> str:
        _, ext = os.path.splitext(file_path)

        try:
            if ext.lower() == ".txt":
                return self._read_text_file(file_path)

            elif ext.lower() == ".pdf":
                return self._read_pdf(file_path)

            elif ext.lower() == ".docx":
                return self._read_docx(file_path)

            elif ext.lower() == ".pptx":
                return self._read_ppt(file_path)

        except Exception as e:
            logger.error(f"[DOC] Error reading {file_path}: {e}")
            return ""

    def _read_text_file(self, path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except:
            with open(path, "r", encoding="latin-1") as f:
                return f.read()

    def _read_pdf(self, path):
        try:
            return extract_text(path)
        except Exception as e:
            logger.error(f"[DOC] PDF read error: {e}")
            return ""

    def _read_docx(self, path):
        try:
            doc = Document(path)
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            logger.error(f"[DOC] DOCX read error: {e}")
            return ""

    def _read_ppt(self, path):
        try:
            prs = Presentation(path)
            text_content = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_content.append(shape.text)

                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text
                    text_content.append(notes)

            return "\n".join(text_content)

        except Exception as e:
            logger.error(f"[DOC] PPTX read error: {e}")
            return ""

    # Text Chunking
    def clean_text(self, text: str) -> str:
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    def chunk_text(self, text: str) -> List[str]:
        text = self.clean_text(text)
        if not text:
            return []

        chunks = []
        start = 0

        while start < len(text):
            end = start + self.chunk_size
            chunk = text[start:end]

            # Smart cut on sentence boundary
            if end < len(text):
                last_period = chunk.rfind(".")
                if last_period != -1:
                    chunk = chunk[:last_period + 1]

            chunks.append(chunk.strip())

            start += self.chunk_size - self.overlap

        return chunks

    # Process Course Documents
    def process_course_documents(self, course_id: int) -> List[str]:
        logger.info(f"[DOC] Processing course {course_id}")

        folder = os.path.join(self.base_courses_dir, str(course_id), "documents")
        chunks = []

        if not os.path.exists(folder):
            logger.warning(f"[DOC] Folder not found: {folder}")
            return chunks

        files = os.listdir(folder)

        for filename in files:
            file_path = os.path.join(folder, filename)
            _, ext = os.path.splitext(filename)

            if ext.lower() not in self.supported_extensions:
                logger.warning(f"[DOC] Unsupported type: {filename}")
                continue

            text = self.extract_text_from_file(file_path)
            if not text:
                continue

            file_chunks = self.chunk_text(text)
            chunks.extend(file_chunks)

            logger.info(f"[DOC] {filename} → {len(file_chunks)} chunks")

        logger.info(f"[DOC] Total chunks: {len(chunks)}")
        return chunks


# Global instance
document_processor = DocumentProcessor()

